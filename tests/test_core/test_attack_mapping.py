"""Tests for MITRE ATT&CK mapping."""

from __future__ import annotations

from openpyxl import load_workbook

from deeplook.core.attack_mapping import (
    ENTITY_TYPE_TECHNIQUES,
    PLUGIN_TECHNIQUE_MAP,
    TECHNIQUES,
    annotate_entities,
    map_entities,
    map_workspace,
    techniques_for_entity,
    techniques_for_plugin,
)
from deeplook.core.entity import EntityType, create_entity
from deeplook.core.workspace import Workspace
from deeplook.output.excel_export import export_xlsx


def test_all_mapped_ids_exist_in_registry():
    for plugin, ids in PLUGIN_TECHNIQUE_MAP.items():
        for tid in ids:
            assert tid in TECHNIQUES, f"{plugin} references unknown technique {tid}"
    for etype, ids in ENTITY_TYPE_TECHNIQUES.items():
        for tid in ids:
            assert tid in TECHNIQUES, f"{etype} references unknown technique {tid}"


def test_technique_urls_are_wellformed():
    assert TECHNIQUES["T1595"].url == "https://attack.mitre.org/techniques/T1595/"
    assert TECHNIQUES["T1595.001"].url == "https://attack.mitre.org/techniques/T1595/001/"


def test_techniques_for_plugin():
    assert "T1596.002" in techniques_for_plugin("domain-whois")
    assert techniques_for_plugin("does-not-exist") == []


def test_techniques_for_entity_via_provenance():
    entity = create_entity("example.com", EntityType.DOMAIN)
    entity.add_provenance("domain-whois", 0.9)
    ids = techniques_for_entity(entity)
    assert "T1596.002" in ids
    assert "T1590.001" in ids


def test_techniques_for_entity_type_fallback():
    entity = create_entity("user@example.com", EntityType.EMAIL)
    assert techniques_for_entity(entity) == ["T1589.002"]


def test_map_entities_aggregates_coverage():
    domain = create_entity("example.com", EntityType.DOMAIN)
    domain.add_provenance("domain-whois", 0.9)
    ip = create_entity("8.8.8.8", EntityType.IP)
    ip.add_provenance("ip-geoloc", 0.8)

    report = map_entities([domain, ip])

    assert report.entity_count == 2
    assert "Reconnaissance" in report.tactics
    ids = {u.technique.id for u in report.techniques}
    assert {"T1596.002", "T1590.005"} <= ids


def test_annotate_entities_persists_metadata(tmp_path):
    ws = Workspace("attack-annotate", base_dir=tmp_path)
    entity = create_entity("example.com", EntityType.DOMAIN)
    entity.add_provenance("domain-whois", 0.9)
    ws.add_entity(entity)

    annotate_entities(ws)

    reloaded = ws.find_entity_by_value("example.com")
    assert "T1596.002" in reloaded.metadata["attack"]


def test_map_workspace_uses_all_entities(tmp_path):
    ws = Workspace("attack-ws", base_dir=tmp_path)
    ws.add_entity(create_entity("example.com", EntityType.DOMAIN))
    ws.add_entity(create_entity("8.8.8.8", EntityType.IP))

    report = map_workspace(ws)

    assert report.entity_count == 2
    assert len(report.techniques) >= 1


def test_excel_export_includes_attack_sheet(tmp_path):
    ws = Workspace("attack-export", base_dir=tmp_path)
    entity = create_entity("example.com", EntityType.DOMAIN)
    entity.add_provenance("domain-whois", 0.9)
    ws.add_entity(entity)

    out = export_xlsx(ws, tmp_path / "report.xlsx", report="attack")

    wb = load_workbook(out)
    assert "ATT&CK" in wb.sheetnames
    attack_ws = wb["ATT&CK"]
    assert attack_ws.cell(row=1, column=1).value == "Technique ID"
    wb.close()
