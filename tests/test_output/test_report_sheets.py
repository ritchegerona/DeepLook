"""Tests for report templates and the new monitoring export sheets."""

from __future__ import annotations

from openpyxl import load_workbook

from deeplook.core.entity import EntityType, create_entity
from deeplook.core.workspace import Workspace
from deeplook.output.excel_export import export_xlsx
from deeplook.output.report_templates import get_template, list_templates


def _workspace(tmp_path):
    ws = Workspace("report-sheets", base_dir=tmp_path)

    ingested = create_entity("evil.example", EntityType.DOMAIN)
    ingested.add_provenance("ingest:feed.txt", 0.7)
    ingested.add_tag("ingested")
    ws.add_entity(ingested)

    ws.add_entity(create_entity("8.8.8.8", EntityType.IP))
    ws.add_watch("domain", "evil.example", tags=["campaign-a"], note="watch this")
    return ws


def test_monitoring_report_has_watchlist_and_ingestion(tmp_path):
    ws = _workspace(tmp_path)

    out = export_xlsx(ws, tmp_path / "monitoring.xlsx", report="monitoring")

    wb = load_workbook(out)
    assert {"Summary", "Watchlist", "Ingestion", "IOCs", "ATT&CK", "Timeline"} <= set(wb.sheetnames)
    wb.close()


def test_watchlist_sheet_contents(tmp_path):
    ws = _workspace(tmp_path)

    out = export_xlsx(ws, tmp_path / "watchlist.xlsx", report="watchlist")

    wb = load_workbook(out)
    assert wb.sheetnames == ["Watchlist"]
    sheet = wb["Watchlist"]
    assert sheet.cell(row=1, column=3).value == "Value"
    assert sheet.cell(row=2, column=3).value == "evil.example"
    assert "campaign-a" in sheet.cell(row=2, column=4).value
    wb.close()


def test_ingestion_sheet_lists_feed(tmp_path):
    ws = _workspace(tmp_path)

    out = export_xlsx(ws, tmp_path / "ingest.xlsx", report="full")

    wb = load_workbook(out)
    assert "Watchlist" in wb.sheetnames
    assert "Ingestion" in wb.sheetnames
    sheet = wb["Ingestion"]
    assert sheet.cell(row=2, column=2).value == "evil.example"
    assert sheet.cell(row=2, column=3).value == "feed.txt"
    wb.close()


def test_summary_counts_watch_and_ingest(tmp_path):
    ws = _workspace(tmp_path)

    out = export_xlsx(ws, tmp_path / "summary.xlsx", report="executive")

    wb = load_workbook(out)
    summary = wb["Summary"]
    metrics = {summary.cell(row=r, column=1).value: summary.cell(row=r, column=2).value for r in range(2, 10)}
    assert metrics["Watched Indicators"] == 1
    assert metrics["Ingested Indicators"] == 1
    wb.close()


def test_monitoring_template_registered():
    names = {t["name"] for t in list_templates()}
    assert "monitoring" in names

    template = get_template("monitoring")
    assert "Watchlist" in template.sheets
    assert "Ingestion" in template.sheets


def test_technical_template_includes_monitoring_sheets():
    sheets = get_template("technical").sheets
    assert "Watchlist" in sheets
    assert "Ingestion" in sheets
