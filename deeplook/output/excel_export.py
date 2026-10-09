"""Excel workbook export for DeepLook workspaces."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from deeplook.core.attack_mapping import map_workspace
from deeplook.core.entity import EntityType
from deeplook.core.workspace import Workspace

_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

_ENTITY_TYPE_FILLS = {
    EntityType.DOMAIN: PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    EntityType.IP: PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
    EntityType.EMAIL: PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    EntityType.USERNAME: PatternFill(start_color="D9B3FF", end_color="D9B3FF", fill_type="solid"),
    EntityType.SUBDOMAIN: PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
}


def _style_header_row(ws) -> None:
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = _THIN_BORDER


def _auto_filter_and_freeze(ws) -> None:
    if ws.max_row > 0 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"


def _auto_size_columns(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, values_only=False):
            for cell in row:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_length + 4, 50)


def _apply_confidence_conditional(ws, col_letter: str) -> None:
    for row_idx in range(2, ws.max_row + 1):
        cell = ws[f"{col_letter}{row_idx}"]
        try:
            value = float(cell.value) if cell.value else 0
        except (ValueError, TypeError):
            continue
        if value > 80:
            cell.fill = _GREEN_FILL
        elif value >= 50:
            cell.fill = _YELLOW_FILL
        else:
            cell.fill = _RED_FILL


def _apply_entity_type_coloring(ws, col_letter: str) -> None:
    for row_idx in range(2, ws.max_row + 1):
        cell = ws[f"{col_letter}{row_idx}"]
        if cell.value:
            try:
                etype = EntityType(cell.value)
                if etype in _ENTITY_TYPE_FILLS:
                    cell.fill = _ENTITY_TYPE_FILLS[etype]
            except (ValueError, KeyError):
                pass


def _write_row(ws, row_idx: int, values: list) -> None:
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _build_summary_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.active
    ws.title = "Summary"
    entities = ws_obj.all_entities()
    edges = ws_obj.all_edges()
    notes = ws_obj.get_notes()

    summary = {
        "Workspace": ws_obj.name,
        "Total Entities": len(entities),
        "Total Relationships": len(edges),
        "Total Notes": len(notes),
    }

    type_counts: dict[str, int] = {}
    for e in entities:
        type_counts[e.type.value] = type_counts.get(e.type.value, 0) + 1

    ws.append(["Metric", "Value"])
    _style_header_row(ws)
    for k, v in summary.items():
        ws.append([k, v])
    ws.append([])
    ws.append(["Entity Type", "Count"])
    for etype, count in sorted(type_counts.items()):
        ws.append([etype, count])

    _auto_size_columns(ws)


def _build_entities_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Entities")
    ws.append(["ID", "Type", "Value", "Confidence", "First Seen", "Last Seen", "Tags", "Notes"])
    _style_header_row(ws)

    for i, entity in enumerate(ws_obj.all_entities(), 2):
        _write_row(
            ws,
            i,
            [
                entity.id,
                entity.type.value,
                entity.value,
                entity.confidence,
                entity.first_seen.isoformat(),
                entity.last_seen.isoformat(),
                ", ".join(entity.tags),
                "; ".join(entity.notes),
            ],
        )

    _apply_entity_type_coloring(ws, "B")
    _apply_confidence_conditional(ws, "D")
    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_relationships_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Relationships")
    ws.append(
        [
            "Source ID",
            "Target ID",
            "Relationship",
            "Confidence",
            "Plugin",
            "First Seen",
            "Last Seen",
        ]
    )
    _style_header_row(ws)

    entity_map = {e.id: e for e in ws_obj.all_entities()}
    for i, edge in enumerate(ws_obj.all_edges(), 2):
        src = entity_map.get(edge.source_id)
        tgt = entity_map.get(edge.target_id)
        _write_row(
            ws,
            i,
            [
                f"[{src.type.value}] {src.value}" if src else edge.source_id,
                f"[{tgt.type.value}] {tgt.value}" if tgt else edge.target_id,
                edge.label,
                edge.confidence,
                edge.source_plugin,
                edge.first_seen.isoformat(),
                edge.last_seen.isoformat(),
            ],
        )

    _apply_confidence_conditional(ws, "D")
    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_domain_details_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Domain Details")
    ws.append(["Domain", "Registrar", "Created", "Expires", "Name Servers", "Confidence"])
    _style_header_row(ws)

    for i, entity in enumerate(ws_obj.entities_by_type(EntityType.DOMAIN), 2):
        meta = entity.metadata
        _write_row(
            ws,
            i,
            [
                entity.value,
                meta.get("registrar", ""),
                meta.get("creation_date", ""),
                meta.get("expiration_date", ""),
                ", ".join(meta.get("name_servers", [])),
                entity.confidence,
            ],
        )

    _apply_confidence_conditional(ws, "F")
    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_ip_details_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("IP Details")
    ws.append(["IP", "ASN", "Organization", "Country", "ISP", "Confidence"])
    _style_header_row(ws)

    for i, entity in enumerate(ws_obj.entities_by_type(EntityType.IP), 2):
        meta = entity.metadata
        _write_row(
            ws,
            i,
            [
                entity.value,
                meta.get("asn", ""),
                meta.get("org", ""),
                meta.get("country", ""),
                meta.get("isp", ""),
                entity.confidence,
            ],
        )

    _apply_confidence_conditional(ws, "F")
    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_username_results_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Username Results")
    ws.append(["Username", "Platform", "URL", "Status", "Confidence"])
    _style_header_row(ws)

    for i, entity in enumerate(ws_obj.entities_by_type(EntityType.USERNAME), 2):
        profiles = entity.metadata.get("profiles", [])
        if profiles:
            for profile in profiles:
                _write_row(
                    ws,
                    i,
                    [
                        entity.value,
                        profile.get("platform", ""),
                        profile.get("url", ""),
                        profile.get("status", ""),
                        entity.confidence,
                    ],
                )
                i += 1
        else:
            _write_row(ws, i, [entity.value, "", "", "", entity.confidence])
            i += 1

    _apply_confidence_conditional(ws, "E")
    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_email_analysis_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Email Analysis")
    ws.append(["Email", "Valid", "Disposable", "Provider", "Confidence"])
    _style_header_row(ws)

    for i, entity in enumerate(ws_obj.entities_by_type(EntityType.EMAIL), 2):
        meta = entity.metadata
        _write_row(
            ws,
            i,
            [
                entity.value,
                meta.get("valid", ""),
                meta.get("disposable", ""),
                meta.get("provider", ""),
                entity.confidence,
            ],
        )

    _apply_confidence_conditional(ws, "E")
    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_iocs_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("IOCs")
    ws.append(["Type", "Value", "Confidence", "First Seen", "Tags"])
    _style_header_row(ws)

    ioc_types = {
        EntityType.IP,
        EntityType.DOMAIN,
        EntityType.HASH,
        EntityType.URL,
        EntityType.EMAIL,
    }
    iocs = [e for e in ws_obj.all_entities() if e.type in ioc_types]
    for i, entity in enumerate(iocs, 2):
        _write_row(
            ws,
            i,
            [
                entity.type.value,
                entity.value,
                entity.confidence,
                entity.first_seen.isoformat(),
                ", ".join(entity.tags),
            ],
        )

    _apply_entity_type_coloring(ws, "A")
    _apply_confidence_conditional(ws, "C")
    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_correlation_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Correlation")
    ws.append(["Type", "Finding", "Entities Involved"])
    _style_header_row(ws)

    edges = ws_obj.all_edges()
    entity_map = {e.id: e for e in ws_obj.all_entities()}

    ip_domains: dict[str, list[str]] = {}
    for edge in edges:
        if edge.relationship.value == "resolves_to":
            src = entity_map.get(edge.source_id)
            tgt = entity_map.get(edge.target_id)
            if src and tgt:
                ip_domains.setdefault(tgt.value, []).append(src.value)

    row_idx = 2
    for ip, domains in ip_domains.items():
        if len(domains) > 1:
            _write_row(
                ws,
                row_idx,
                [
                    "Shared Infrastructure",
                    f"IP {ip} resolves from {len(domains)} domains",
                    ", ".join(domains),
                ],
            )
            row_idx += 1

    email_entities = ws_obj.entities_by_type(EntityType.EMAIL)
    email_domains: dict[str, list[str]] = {}
    for email in email_entities:
        parts = email.value.split("@")
        if len(parts) == 2:
            related = [e for e in ws_obj.all_entities() if e.type == EntityType.DOMAIN]
            for domain in related:
                if domain.value == parts[1]:
                    email_domains.setdefault(email.value, []).append(domain.value)

    for email_val, domains in email_domains.items():
        if len(domains) > 1:
            _write_row(
                ws,
                row_idx,
                [
                    "Shared Contact",
                    f"Email {email_val} associated with {len(domains)} domains",
                    ", ".join(domains),
                ],
            )
            row_idx += 1

    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_attack_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("ATT&CK")
    ws.append(["Technique ID", "Technique", "Tactic", "Entities", "URL"])
    _style_header_row(ws)

    report = map_workspace(ws_obj)
    for i, usage in enumerate(report.techniques, 2):
        _write_row(
            ws,
            i,
            [
                usage.technique.id,
                usage.technique.name,
                usage.technique.tactic,
                ", ".join(usage.entities),
                usage.technique.url,
            ],
        )

    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_timeline_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Timeline")
    ws.append(["Timestamp", "Event", "Target", "Details"])
    _style_header_row(ws)

    for i, event in enumerate(ws_obj.get_timeline(), 2):
        _write_row(
            ws,
            i,
            [
                event.get("timestamp", ""),
                event.get("event", ""),
                event.get("target", ""),
                event.get("details", ""),
            ],
        )

    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


def _build_notes_sheet(wb: Workbook, ws_obj: Workspace) -> None:
    ws = wb.create_sheet("Notes")
    ws.append(["ID", "Entity ID", "Content", "Created At"])
    _style_header_row(ws)

    for i, note in enumerate(ws_obj.get_notes(), 2):
        _write_row(
            ws,
            i,
            [
                note.get("id", ""),
                note.get("entity_id", ""),
                note.get("content", ""),
                note.get("created_at", ""),
            ],
        )

    _auto_filter_and_freeze(ws)
    _auto_size_columns(ws)


_SHEET_BUILDERS = {
    "Summary": _build_summary_sheet,
    "Entities": _build_entities_sheet,
    "Relationships": _build_relationships_sheet,
    "Domain Details": _build_domain_details_sheet,
    "IP Details": _build_ip_details_sheet,
    "Username Results": _build_username_results_sheet,
    "Email Analysis": _build_email_analysis_sheet,
    "IOCs": _build_iocs_sheet,
    "Correlation": _build_correlation_sheet,
    "ATT&CK": _build_attack_sheet,
    "Timeline": _build_timeline_sheet,
    "Notes": _build_notes_sheet,
}

_FULL_REPORT_SHEETS = [
    "Summary",
    "Entities",
    "Relationships",
    "Domain Details",
    "IP Details",
    "Username Results",
    "Email Analysis",
    "IOCs",
    "Correlation",
    "ATT&CK",
    "Timeline",
    "Notes",
]

_REPORT_SHEET_MAP: dict[str, list[str]] = {
    "full": _FULL_REPORT_SHEETS,
    "executive": ["Summary", "IOCs"],
    "technical": _FULL_REPORT_SHEETS,
    "ioc": ["IOCs"],
    "attack": ["Summary", "ATT&CK"],
}


def export_xlsx(workspace: Workspace, output_path: Path, report: str = "full") -> Path:
    wb = Workbook()
    sheets = _REPORT_SHEET_MAP.get(report, _FULL_REPORT_SHEETS)

    for sheet_name in sheets:
        builder = _SHEET_BUILDERS.get(sheet_name)
        if builder:
            if sheet_name == "Summary":
                wb.active.title = sheet_name
            builder(wb, workspace)

    if "Summary" not in sheets and wb.sheetnames:
        del wb[wb.sheetnames[0]]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    wb.close()
    return output_path
